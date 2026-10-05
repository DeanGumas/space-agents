"""The xlsx diff/repair script normally runs inside the Pyodide sandbox; it's plain openpyxl, so test it on the host."""

import json
import subprocess
import sys
from datetime import datetime

import openpyxl

from space_agents.tools.registry import XLSX_DIFF


def run_diff(tmp_path, old_wb, new_wb):
    (tmp_path / "in").mkdir()
    (tmp_path / "out").mkdir()
    if old_wb is not None:
        old_wb.save(tmp_path / "in" / "old.xlsx")
    new_wb.save(tmp_path / "in" / "new.xlsx")
    r = subprocess.run([sys.executable, "-c", XLSX_DIFF], cwd=tmp_path, capture_output=True, text=True, check=True)
    return json.loads(r.stdout.strip().splitlines()[-1])


def workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "POs"
    for row in rows:
        ws.append(row)
    ws["B2"].number_format = "yyyy-mm-dd"
    return wb


def test_counts_changes_and_repairs_text_dates(tmp_path):
    old = workbook([["PO", "Due", "Status"], [45821, datetime(2026, 10, 9), "Open"]])
    new = workbook([["PO", "Due", "Status"], [45821, "2026-10-14", "Delayed"]])
    out = run_diff(tmp_path, old, new)
    assert out["changed_cells"] == 2 and out["dates_repaired"] == 1 and out["date_to_text"] == 0
    repaired = openpyxl.load_workbook(tmp_path / "out" / "new.xlsx")["POs"]["B2"]
    assert repaired.value == datetime(2026, 10, 14) and repaired.number_format == "yyyy-mm-dd"


def test_non_iso_text_in_date_cell_is_flagged_not_changed(tmp_path):
    old = workbook([["PO", "Due"], [1, datetime(2026, 10, 9)]])
    new = workbook([["PO", "Due"], [1, "next week"]])
    out = run_diff(tmp_path, old, new)
    assert out["date_to_text"] == 1 and out["dates_repaired"] == 0
    assert not (tmp_path / "out" / "new.xlsx").exists()


def test_new_workbook_without_original(tmp_path):
    out = run_diff(tmp_path, None, workbook([["a"], [1], [2]]))
    assert out["sheets"] == [{"name": "POs", "rows": 3, "status": "new"}]


def test_detects_lost_formatting(tmp_path):
    from openpyxl.styles import PatternFill

    old = workbook([["PO", "Due"], [1, datetime(2026, 10, 9)]])
    old.active["A1"].fill = PatternFill("solid", fgColor="1F3A5F")
    new = openpyxl.Workbook()  # like pandas: same values, no styles
    new.active.title = "POs"
    for row in (["PO", "Due"], [1, datetime(2026, 10, 9)]):
        new.active.append(row)
    out = run_diff(tmp_path, old, new)
    assert out["formats_lost"] == 1 and out["changed_cells"] == 0  # the header fill (openpyxl auto-formats dates)
