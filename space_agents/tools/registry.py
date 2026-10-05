"""The tools an agent can call, and how each one runs (PLAN.md §2.7).

Every tool returns a ToolOutcome: `result` goes back to the model, `summary` is the plain-language step shown to
the user ("Read PO Tracker.xlsx"). Paths are always checked by FolderAccess; files the agent produces go through
ChangeManager as proposed changes.
"""

from __future__ import annotations

import ast
import difflib
import json
import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from space_agents.changes import ChangeManager
from space_agents.documents.extract import UnsupportedDocument
from space_agents.documents.library import DocumentLibrary
from space_agents.manifest import Manifest
from space_agents.sandbox.pyodide import PyodideSandbox, RunResult, SandboxUnavailable
from space_agents.tools.access import AccessDenied, FolderAccess, display_path

log = logging.getLogger(__name__)

MAX_LIST_ENTRIES = 300
SANDBOX_PACKAGES = "pandas, numpy, openpyxl, python-docx (import docx), pypdf, xlsxwriter, matplotlib"

TOOL_SPECS: list[dict[str, Any]] = [
    {
        "type": "function",
        "function": {
            "name": "list_files",
            "description": "List files in the agent's folders (recursive). Use this to find files before reading them.",
            "parameters": {
                "type": "object",
                "properties": {
                    "folder": {"type": "string", "description": "Folder to list. Omit to list all granted folders."},
                    "pattern": {"type": "string", "description": "Optional glob filter, e.g. '*.pdf' or '*PO*'."},
                },
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_document",
            "description": (
                "Read the text of a file: PDF, Word (.docx), Excel (.xlsx), PowerPoint (.pptx), or text. "
                "Large documents return a page range at a time; the header says how many pages exist."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "File path (as shown by list_files)."},
                    "pages": {"type": "string", "description": "Optional page range like '1-5' or '2,7,10-12'."},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "search_document",
            "description": "Find the pages of a (large) document most relevant to a query. Returns page numbers.",
            "parameters": {
                "type": "object",
                "properties": {"path": {"type": "string"}, "query": {"type": "string"}},
                "required": ["path", "query"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "run_python",
            "description": (
                "Run Python 3 in a secure sandbox to analyze, create, or update files. "
                f"Available packages: {SANDBOX_PACKAGES}. No network access. "
                "Files listed in input_files are copied to ./in/<filename>. To create or update a file, save it to "
                "./out/<filename>; using an input's filename updates that file, other names create new files in "
                "output_folder (or next to the first input). Example: wb = openpyxl.load_workbook('in/Report.xlsx'); "
                "...; wb.save('out/Report.xlsx'). Match cell types: compare numbers as numbers, write dates as "
                "datetime.date. To edit an existing Excel file use openpyxl (pandas rewrites the whole file and "
                "drops its formatting). The user reviews changes before they're applied. print() what you need."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "description": {
                        "type": "string",
                        "description": "Short plain-English summary of what this does, shown to the user.",
                    },
                    "code": {"type": "string", "description": "Python code to run."},
                    "input_files": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "Paths of files to copy into ./in/.",
                    },
                    "output_folder": {
                        "type": "string",
                        "description": "Folder for new output files (must be a read & write folder).",
                    },
                },
                "required": ["description", "code"],
            },
        },
    },
]

XLSX_DIFF = r"""
import datetime, json, openpyxl
from openpyxl.utils import get_column_letter

def show(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime) and v.time() == datetime.time(0):
        return v.date().isoformat()
    return str(v)[:60]

def load(p):
    try:
        return openpyxl.load_workbook(p)
    except FileNotFoundError:
        return None

old, new = load("in/old.xlsx"), load("in/new.xlsx")
import re
out = {"sheets": [], "samples": [], "changed_cells": 0, "date_to_text": 0, "date_to_text_cell": "", "dates_repaired": 0,
       "formats_lost": 0}
ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?$")
for ws in new.worksheets:
    entry = {"name": ws.title, "rows": ws.max_row}
    o = old[ws.title] if old is not None and ws.title in old.sheetnames else None
    if o is None:
        entry["status"] = "new" if old is None else "added"
    else:
        entry["old_rows"] = o.max_row
        rows, cols = max(ws.max_row, o.max_row), max(ws.max_column, o.max_column)
        changed = 0
        if rows * cols <= 300_000:
            for r in range(1, rows + 1):
                for c in range(1, cols + 1):
                    a, b = o.cell(r, c).value, ws.cell(r, c).value
                    # Guardrail: a date cell overwritten with an ISO date string becomes a real date again,
                    # keeping the original number format (small models often write "2026-10-14" as text).
                    m = ISO.match(b.strip()) if isinstance(b, str) else None
                    if isinstance(a, (datetime.date, datetime.datetime)) and m:
                        y, mo, d, hh, mm, ss = m.groups()
                        b = datetime.datetime(int(y), int(mo), int(d), int(hh or 0), int(mm or 0), int(ss or 0))
                        ws.cell(r, c).value = b
                        ws.cell(r, c).number_format = o.cell(r, c).number_format
                        out["dates_repaired"] += 1
                    oc, nc = o.cell(r, c), ws.cell(r, c)
                    if (oc.number_format != "General" and nc.number_format == "General") or (
                        oc.fill.fill_type and not nc.fill.fill_type
                    ):
                        out["formats_lost"] += 1
                    if a != b:
                        changed += 1
                        if isinstance(a, (datetime.date, datetime.datetime)) and isinstance(b, str):
                            out["date_to_text"] += 1
                            out["date_to_text_cell"] = out["date_to_text_cell"] or f"{get_column_letter(c)}{r}"
                        if len(out["samples"]) < 12:
                            out["samples"].append({"sheet": ws.title, "cell": f"{get_column_letter(c)}{r}",
                                                   "old": show(a), "new": show(b)})
        entry["changed_cells"] = changed
        out["changed_cells"] += changed
    out["sheets"].append(entry)
if old is not None:
    for name in old.sheetnames:
        if name not in new.sheetnames:
            out["sheets"].append({"name": name, "status": "removed"})
if out["dates_repaired"]:
    new.save("out/new.xlsx")
print(json.dumps(out))
"""


@dataclass
class ToolOutcome:
    result: str
    summary: str
    ok: bool = True
    changes: list[dict[str, Any]] = field(default_factory=list)


class ToolBox:
    def __init__(
        self,
        *,
        agent: Manifest,
        message_id: int,
        docs: DocumentLibrary,
        sandbox: PyodideSandbox,
        changes: ChangeManager,
        protected: list[Path],
    ):
        self.agent = agent
        self.message_id = message_id
        self.docs = docs
        self.sandbox = sandbox
        self.changes = changes
        self.access = FolderAccess(agent.filesystem, protected)

    def specs(self) -> list[dict[str, Any]]:
        return TOOL_SPECS

    async def call(self, name: str, args: dict[str, Any], tool_call_id: int | None = None) -> ToolOutcome:
        try:
            if name == "list_files":
                return self._list_files(str(args.get("folder") or ""), str(args.get("pattern") or "*"))
            if name == "read_document":
                return await self._read(str(args.get("path") or ""), args.get("pages"))
            if name == "search_document":
                return await self._search(str(args.get("path") or ""), str(args.get("query") or ""))
            if name == "run_python":
                return await self._run_python(args, tool_call_id)
            return ToolOutcome(f"Unknown tool {name!r}.", f"Unknown tool {name}", ok=False)
        except (AccessDenied, UnsupportedDocument, FileNotFoundError, SandboxUnavailable, ValueError) as e:
            msg = str(e) if not isinstance(e, FileNotFoundError) else f"File not found: {e.filename or e}"
            return ToolOutcome(f"Error: {msg}", _failed_summary(name, args), ok=False)
        except Exception:  # a tool bug must not end the whole run; the model can still answer
            log.exception("Tool %s failed", name)
            return ToolOutcome("Error: the tool failed unexpectedly.", _failed_summary(name, args), ok=False)

    # --- files and documents ----------------------------------------------------------------------------------

    def _list_files(self, folder: str, pattern: str) -> ToolOutcome:
        bases = [self.access.resolve(folder)] if folder else [r.path for r in self.access.roots]
        lines, count = [], 0
        for base in bases:
            if not base.is_dir():
                raise ValueError(f"{display_path(base)} is not a folder.")
            lines.append(f"{display_path(base)}/")
            for p in sorted(base.rglob(pattern)):
                rel = p.relative_to(base)
                if any(part.startswith(".") for part in rel.parts) or not p.is_file():
                    continue
                try:
                    self.access.resolve(str(p))
                except AccessDenied:
                    continue
                st = p.stat()
                modified = datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M")
                lines.append(f"  {rel.as_posix()}  ({_size(st.st_size)}, modified {modified})")
                count += 1
                if count >= MAX_LIST_ENTRIES:
                    lines.append(f"  … stopped at {MAX_LIST_ENTRIES} files; use a pattern to narrow the list.")
                    break
        where = folder or "the agent's folders"
        return ToolOutcome(
            "\n".join(lines) if count else "\n".join(lines + ["  (no files)"]), f"Listed files in {where}"
        )

    async def _read(self, raw: str, pages: Any) -> ToolOutcome:
        path = self.access.resolve(raw)
        text = await self.docs.read(path, path.name, str(pages) if pages else None)
        where = f" (pages {pages})" if pages else ""
        return ToolOutcome(text, f"Read {path.name}{where}")

    async def _search(self, raw: str, query: str) -> ToolOutcome:
        path = self.access.resolve(raw)
        return ToolOutcome(await self.docs.search(path, path.name, query), f"Searched {path.name} for “{query}”")

    # --- sandbox ----------------------------------------------------------------------------------------------

    async def _run_python(self, args: dict[str, Any], tool_call_id: int | None) -> ToolOutcome:
        description = str(args.get("description") or "Ran Python").strip()[:200]
        code = str(args.get("code") or "")
        if not code.strip():
            raise ValueError("No code was provided.")
        raw_inputs = args.get("input_files") or []
        if isinstance(raw_inputs, str):
            raw_inputs = [raw_inputs]
        inputs: dict[str, Path] = {}
        for raw in raw_inputs:
            path = self.access.resolve(str(raw))
            if not path.is_file():
                raise ValueError(f"{display_path(path)} is not a file.")
            if path.name in inputs and inputs[path.name] != path:
                raise ValueError(f"Two input files are named {path.name}; pass them in separate runs.")
            inputs[path.name] = path
        aliases = self._path_literals(code, inputs)
        # If the agent already proposed a change to an input in this message, it works on that latest version.
        sources = {}
        for name, path in inputs.items():
            pending = self.changes.pending_for(self.message_id, path)
            sources[name] = Path(pending["staged"]) if pending and Path(pending["staged"]).exists() else path

        result = await self.sandbox.run(code, sources, aliases=aliases)
        try:
            lines = [f"{'Succeeded' if result.ok else 'Failed'} in {result.duration_s:.1f}s."]
            if result.stdout.strip():
                lines.append("Output:\n" + result.stdout.strip())
            if result.stderr.strip():
                lines.append("Warnings/stderr:\n" + result.stderr.strip()[-3000:])
            if result.error:
                lines.append("Error:\n" + result.error)
                if "No such file or directory" in result.error:
                    lines.append(
                        "Hint: pass the file in input_files and open it as 'in/<filename>'; save results to "
                        "'out/<filename>'."
                    )
            proposed = []
            if result.ok and result.outputs:
                lines.append("Files:")
                for out in result.outputs:
                    lines.append("- " + await self._propose(out, inputs, args, description, tool_call_id, proposed))
            elif result.outputs:
                lines.append("Output files were not saved because the code failed.")
            return ToolOutcome("\n".join(lines), description, ok=result.ok, changes=proposed)
        finally:
            self.sandbox.release(result)

    def _path_literals(self, code: str, inputs: dict[str, Path]) -> dict[str, str]:
        """Find string literals in the code that name files in the agent's folders; add them as inputs.

        Returns {literal: input name}, so the sandbox can place a copy exactly where the code looks for it. Small
        models often open files by their real path instead of ./in/<name>; this makes that just work (writes still
        come back as proposed changes). The code is only parsed here, never executed on the host.
        """
        try:
            tree = ast.parse(code)
        except SyntaxError:
            return {}  # the sandbox will report the syntax error
        aliases: dict[str, str] = {}
        for node in ast.walk(tree):
            if not (isinstance(node, ast.Constant) and isinstance(node.value, str)):
                continue
            literal = node.value.strip()
            if not literal or len(literal) > 500 or "\n" in literal:
                continue
            lookup = literal.removeprefix("./")
            prefix = lookup.split("/", 1)[0] if lookup.startswith(("in/", "out/")) else ""
            if prefix:  # the sandbox convention: in/<name> or out/<name> → the file of that name in the folders
                lookup = lookup.split("/", 1)[1]
            try:
                path = self.access.resolve(lookup)
            except (AccessDenied, ValueError, OSError):
                continue
            if not path.is_file() or (path.name in inputs and inputs[path.name] != path):
                continue
            inputs[path.name] = path
            if prefix != "in":  # in/<name> is where inputs land anyway
                aliases[literal.removeprefix("./")] = path.name
        return aliases

    async def _propose(
        self,
        out: Path,
        inputs: dict[str, Path],
        args: dict[str, Any],
        description: str,
        tool_call_id: int | None,
        proposed: list[dict[str, Any]],
    ) -> str:
        try:
            if out.name in inputs:
                target = self.access.resolve(str(inputs[out.name]), write=True)
            else:
                folder = str(args.get("output_folder") or "")
                if folder and not self.access.resolve(folder).is_dir():
                    folder = ""  # e.g. the model confused the sandbox's out/ with a real folder
                if folder:
                    base = self.access.resolve(folder, write=True)
                elif inputs:
                    base = next(iter(inputs.values())).parent
                else:
                    writable = [r for r in self.access.roots if r.writable]
                    if len(writable) != 1:
                        raise AccessDenied("Specify output_folder (a read & write folder) for new files.")
                    base = writable[0].path
                target = self.access.resolve(str(base / out.name), write=True)
        except AccessDenied as e:
            return f"{out.name}: not saved. {e}"
        details = await self._summarize(out, target)
        if target.exists() and _unchanged(out, target, details):
            return (
                f"{target.name}: no changes. The saved file has the same content as the original, so nothing was "
                "proposed. Check your matching logic (e.g. numbers vs. text) if you expected changes."
            )
        summary = {"description": description, **details}
        change = self.changes.propose(
            agent_id=self.agent.id,
            message_id=self.message_id,
            tool_call_id=tool_call_id,
            target=target,
            source=out,
            summary=summary,
            policy=self.agent.file_changes,
        )
        proposed.append(change)
        verb = "Updated" if change["kind"] == "modify" else "Created"
        state = {
            "applied": "saved",
            "pending": "waiting for the user to review and apply it",
            "conflicted": "the file changed in the meantime; the user must decide",
        }.get(change["status"], change["status"])
        return f"{verb} {change['target']}: {state}.{_describe(details)}"

    async def _summarize(self, new: Path, target: Path) -> dict[str, Any]:
        suffix = new.suffix.lower()
        details: dict[str, Any] = {"size": new.stat().st_size}
        if suffix in (".xlsx", ".xlsm"):
            inputs = {"new.xlsx": new}
            if target.exists():
                inputs["old.xlsx"] = target
            r: RunResult = await self.sandbox.run(XLSX_DIFF, inputs, timeout=60)
            try:
                if r.ok and r.stdout.strip():
                    details["xlsx"] = json.loads(r.stdout.strip().splitlines()[-1])
                repaired = next((o for o in r.outputs if o.name == "new.xlsx"), None)
                if repaired is not None:
                    new.write_bytes(repaired.read_bytes())
            except json.JSONDecodeError:
                log.warning("Couldn't parse xlsx diff output")
            finally:
                self.sandbox.release(r)
        elif suffix in (".txt", ".md", ".csv", ".tsv", ".json") and target.exists():
            old = target.read_text(encoding="utf-8", errors="replace").splitlines()
            cur = new.read_text(encoding="utf-8", errors="replace").splitlines()
            diff = list(difflib.unified_diff(old, cur, lineterm="", n=1))
            details["text_diff"] = "\n".join(diff[2:80])
            details["lines_added"] = sum(1 for d in diff if d.startswith("+") and not d.startswith("+++"))
            details["lines_removed"] = sum(1 for d in diff if d.startswith("-") and not d.startswith("---"))
        return details


def _unchanged(new: Path, old: Path, details: dict[str, Any]) -> bool:
    if new.read_bytes() == old.read_bytes():
        return True
    x = details.get("xlsx")
    return (
        bool(x)
        and x.get("changed_cells", 0) == 0
        and all(s.get("status") is None and s.get("rows") == s.get("old_rows") for s in x["sheets"])
    )


def _describe(details: dict[str, Any]) -> str:
    """Summarize a change for the model, so it can verify its own work."""
    x = details.get("xlsx")
    if x:
        parts = []
        for sh in x["sheets"]:
            if sh.get("status"):
                parts.append(f"sheet '{sh['name']}' {sh['status']}")
            elif sh.get("changed_cells") or sh.get("rows") != sh.get("old_rows"):
                parts.append(
                    f"'{sh['name']}': {sh['changed_cells']} cell(s) changed, rows {sh['old_rows']}→{sh['rows']}"
                )
        samples = "; ".join(f"{c['cell']}: {c['old'] or '(empty)'} → {c['new'] or '(empty)'}" for c in x["samples"][:8])
        text = f" Changes: {', '.join(parts) or 'none'}." + (f" Examples: {samples}." if samples else "")
        if x.get("dates_repaired"):
            text += f" Note: {x['dates_repaired']} date cell(s) written as text were converted back to real dates."
        if x.get("formats_lost"):
            text += (
                f" Warning: formatting was lost on {x['formats_lost']} cell(s) (number formats/fills). If you used "
                "pandas to rewrite the file, redo the edit with openpyxl (load_workbook → change cells → save)."
            )
        if x.get("date_to_text"):
            text += (
                f" Warning: {x['date_to_text']} cell(s) that held dates now hold text (e.g. {x['date_to_text_cell']});"
                " write datetime.date values instead."
            )
        return text
    if "lines_added" in details:
        return f" Changes: +{details['lines_added']} / -{details['lines_removed']} lines."
    return ""


def _size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n} B"


def _failed_summary(name: str, args: dict[str, Any]) -> str:
    target = args.get("path") or args.get("folder") or args.get("description") or ""
    label = {"list_files": "List files", "read_document": "Read", "search_document": "Search", "run_python": "Run"}
    return f"{label.get(name, name)} {Path(str(target)).name if 'path' in args else target}".strip()
